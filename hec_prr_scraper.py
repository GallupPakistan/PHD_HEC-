"""
HEC PRR Dashboard Scraper
=========================
Scrapes University, Discipline, Subject and Year tables from:
https://pcd.hec.gov.pk/#/prr

Each table is paginated independently using its own "Next" button.
All four datasets are written into a single Excel workbook
(Dashboard_Data.xlsx) with four worksheets.

Requirements:
    pip install selenium webdriver-manager pandas openpyxl
"""

import time
import logging
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    StaleElementReferenceException,
    TimeoutException,
    ElementClickInterceptedException,
    NoSuchElementException,
    WebDriverException,
)
from webdriver_manager.chrome import ChromeDriverManager


# =========================================================
# CONFIGURATION
# =========================================================

WEBSITE_URL = "https://pcd.hec.gov.pk/#/prr"

# (name, table_container_xpath, next_button_xpath, output_columns)
SECTIONS = [
    {
        "name": "University",
        "table_xpath": '//*[@id="section3"]/div[1]/div[2]/div/div/div/div',
        "next_xpath": '//*[@id="section3"]/div[1]/div[2]/div/div/div/div/div/div/mat-paginator/div/div/div[2]/button[2]',
        "columns": ["University Name", "Numbers"],
    },
    {
        "name": "Discipline",
        "table_xpath": '//*[@id="section3"]/div[2]/div[2]/div/div/div/div',
        "next_xpath": '//*[@id="section3"]/div[2]/div[2]/div/div/div/div/div/div/mat-paginator/div/div/div[2]/button[2]',
        "columns": ["Discipline", "Numbers"],
    },
    {
        "name": "Subject",
        "table_xpath": '//*[@id="section3"]/div[3]/div[2]/div/div/div/div',
        "next_xpath": '//*[@id="section3"]/div[3]/div[2]/div/div/div/div/div/div/mat-paginator/div/div/div[2]/button[2]',
        "columns": ["Subject", "Numbers"],
    },
    {
        "name": "Year",
        "table_xpath": '//*[@id="section3"]/div[4]/div[2]/div/div/div/div',
        "next_xpath": '//*[@id="section3"]/div[4]/div[2]/div/div/div/div/div/div/mat-paginator/div/div/div[2]/button[2]',
        "columns": ["Year", "Numbers"],
    },
]

OUTPUT_FILE = "Dashboard_Data.xlsx"

PAGE_LOAD_TIMEOUT = 30
ELEMENT_WAIT_TIMEOUT = 20
MAX_RETRIES_PER_PAGE = 3
MAX_EMPTY_NEXT_ATTEMPTS = 3  # safety valve against infinite loops


# =========================================================
# LOGGING
# =========================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(message)s",
    stream=sys.stdout,
)
log = logging.getLogger("hec_scraper")


# =========================================================
# DRIVER SETUP
# =========================================================

def start_driver():
    """Start and configure a Chrome webdriver instance."""
    options = Options()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-notifications")
    # Uncomment the line below to run headless
    # options.add_argument("--headless=new")

    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)
    driver.set_page_load_timeout(PAGE_LOAD_TIMEOUT)
    return driver


# =========================================================
# HELPERS
# =========================================================

def clean_cell(text):
    """Trim whitespace and collapse internal whitespace."""
    if text is None:
        return ""
    return " ".join(text.split()).strip()


def get_table_rows_text(driver, table_xpath):
    """
    Read all visible rows/cells from the given table container.
    Returns a list of tuples representing each row's cell text.
    Uses a defensive approach: collects text from any 'tr' or
    mat-row elements, falling back to generic row/cell divs used
    by Angular Material tables.
    """
    container = driver.find_element(By.XPATH, table_xpath)

    rows_data = []

    # Try standard HTML table rows first
    row_elements = container.find_elements(By.XPATH, ".//tr")
    if not row_elements:
        # Fall back to Angular Material mat-row
        row_elements = container.find_elements(By.XPATH, ".//mat-row")

    for row in row_elements:
        cells = row.find_elements(By.XPATH, ".//td")
        if not cells:
            cells = row.find_elements(By.XPATH, ".//mat-cell")
        if not cells:
            continue

        cell_values = [clean_cell(c.text) for c in cells]

        # Skip fully empty rows
        if not any(cell_values):
            continue

        rows_data.append(tuple(cell_values))

    return rows_data


def wait_for_table(driver, table_xpath):
    """Wait until the table container is visible on the page."""
    WebDriverWait(driver, ELEMENT_WAIT_TIMEOUT).until(
        EC.visibility_of_element_located((By.XPATH, table_xpath))
    )


def get_next_button(driver, next_xpath):
    """Return the Next button element, or None if not found."""
    try:
        return driver.find_element(By.XPATH, next_xpath)
    except NoSuchElementException:
        return None


def is_button_disabled(button):
    """Check whether a mat-paginator Next button is disabled."""
    if button is None:
        return True
    try:
        disabled_attr = button.get_attribute("disabled")
        aria_disabled = button.get_attribute("aria-disabled")
        class_attr = button.get_attribute("class") or ""
        if disabled_attr is not None:
            return True
        if aria_disabled and aria_disabled.lower() == "true":
            return True
        if "mat-button-disabled" in class_attr or "disabled" in class_attr.split():
            return True
        return False
    except StaleElementReferenceException:
        return True


def click_next(driver, next_xpath):
    """
    Click the Next button robustly, handling intercepted clicks
    and stale elements. Returns True if a click was performed.
    """
    for attempt in range(MAX_RETRIES_PER_PAGE):
        try:
            button = get_next_button(driver, next_xpath)
            if button is None or is_button_disabled(button):
                return False

            driver.execute_script(
                "arguments[0].scrollIntoView({block: 'center'});", button
            )
            try:
                button.click()
            except ElementClickInterceptedException:
                # Fall back to a JS click if a normal click is intercepted
                driver.execute_script("arguments[0].click();", button)

            return True

        except StaleElementReferenceException:
            time.sleep(0.5)
            continue
        except WebDriverException:
            time.sleep(0.5)
            continue

    return False


def wait_for_table_refresh(driver, table_xpath, previous_first_row):
    """
    Wait until the table content changes after clicking Next.
    Compares the first row's text before/after to detect refresh.
    Falls back gracefully if comparison cannot be made.
    """

    def _table_changed(drv):
        try:
            rows = get_table_rows_text(drv, table_xpath)
            if not rows:
                return False
            return rows[0] != previous_first_row
        except (StaleElementReferenceException, NoSuchElementException):
            return False

    try:
        WebDriverWait(driver, ELEMENT_WAIT_TIMEOUT).until(_table_changed)
        return True
    except TimeoutException:
        return False


# =========================================================
# CORE SCRAPING LOGIC
# =========================================================

def scrape_table(driver, section):
    """
    Scrape a single section's table across all its pages.
    Returns a deduplicated list of row tuples, preserving order.
    """
    name = section["name"]
    table_xpath = section["table_xpath"]
    next_xpath = section["next_xpath"]

    log.info(f"\nScraping {name}...")

    all_rows = []
    seen_rows = set()
    page_number = 1
    empty_next_attempts = 0

    wait_for_table(driver, table_xpath)

    while True:
        # --- Read current page with retries for stale/transient errors ---
        current_rows = None
        for attempt in range(MAX_RETRIES_PER_PAGE):
            try:
                current_rows = get_table_rows_text(driver, table_xpath)
                break
            except StaleElementReferenceException:
                time.sleep(0.5)
                continue
            except NoSuchElementException:
                time.sleep(0.5)
                continue

        if current_rows is None:
            log.info(f"Page {page_number} : could not read rows, stopping.")
            break

        new_rows_count = 0
        for row in current_rows:
            if row not in seen_rows:
                seen_rows.add(row)
                all_rows.append(row)
                new_rows_count += 1

        log.info(f"Page {page_number} : {len(current_rows)} rows "
                  f"({new_rows_count} new)")

        # Track the first row so we can detect a real page change
        previous_first_row = current_rows[0] if current_rows else None

        # --- Determine if pagination should continue ---
        next_button = get_next_button(driver, next_xpath)
        if is_button_disabled(next_button):
            break

        clicked = click_next(driver, next_xpath)
        if not clicked:
            break

        # Wait for the table to actually refresh with new data
        refreshed = wait_for_table_refresh(driver, table_xpath, previous_first_row)

        if not refreshed:
            empty_next_attempts += 1
            if new_rows_count == 0 or empty_next_attempts >= MAX_EMPTY_NEXT_ATTEMPTS:
                # No new data appeared after clicking Next multiple
                # times in a row -- assume end of data.
                break
            # otherwise give it one more short pause and retry the read
            time.sleep(1)
        else:
            empty_next_attempts = 0

        page_number += 1

    log.info(f"{name} Finished : {len(all_rows)} rows")
    return all_rows


def build_dataframe(rows, columns):
    """
    Convert scraped rows into a cleaned pandas DataFrame:
    trims whitespace (already trimmed), drops fully empty rows,
    removes duplicates while preserving original order, and
    aligns column count with the provided column names.
    """
    if not rows:
        return pd.DataFrame(columns=columns)

    width = len(columns)
    normalized_rows = []
    for row in rows:
        row = list(row)
        if len(row) < width:
            row = row + [""] * (width - len(row))
        elif len(row) > width:
            row = row[:width]
        normalized_rows.append(tuple(row))

    df = pd.DataFrame(normalized_rows, columns=columns)

    # Drop rows where every column is empty
    df = df[~(df.apply(lambda r: all(str(v).strip() == "" for v in r), axis=1))]

    # Remove duplicate rows while preserving first-seen order
    df = df.drop_duplicates(keep="first").reset_index(drop=True)

    return df


# =========================================================
# EXCEL EXPORT
# =========================================================

def save_excel(data_frames, output_file):
    """
    Write each section's DataFrame to its own worksheet in a
    single workbook. Bold headers, freeze the first row, and
    auto-adjust column widths.
    """
    wb = Workbook()
    # Remove the default blank sheet; we'll add our own in order
    default_sheet = wb.active
    wb.remove(default_sheet)

    bold_font = Font(bold=True)

    for sheet_name, df in data_frames.items():
        ws = wb.create_sheet(title=sheet_name)

        # Write header
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = bold_font

        # Write data rows
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Freeze header row
        ws.freeze_panes = "A2"

        # Auto-adjust column widths
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = len(str(col_name))
            for value in df.iloc[:, col_idx - 1]:
                max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    wb.save(output_file)
    log.info("\nExcel Saved.")


# =========================================================
# MAIN
# =========================================================

def main():
    log.info("Opening Website...")

    driver = start_driver()
    data_frames = {}

    try:
        driver.get(WEBSITE_URL)

        # Give the Angular app a moment to bootstrap and wait for the
        # first section's table to appear as a readiness signal.
        try:
            wait_for_table(driver, SECTIONS[0]["table_xpath"])
        except TimeoutException:
            log.info("Initial page load timed out, continuing anyway...")

        for section in SECTIONS:
            try:
                rows = scrape_table(driver, section)
            except Exception as exc:
                log.info(f"Error while scraping {section['name']}: {exc}")
                rows = []

            df = build_dataframe(rows, section["columns"])
            data_frames[section["name"]] = df

        save_excel(data_frames, OUTPUT_FILE)

    finally:
        driver.quit()


if __name__ == "__main__":
    main()
